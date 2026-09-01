#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""render_painel_lancamento — contrato.json → painel HTML (as telas do relatório).

    python3 scripts/render_painel_lancamento.py docs/relatorios/lf64_resultado

Lê SOMENTE o contrato congelado (nunca banco): republicar o painel é re-rodar
este script; preencher as vendas é re-rodar o relatorio_lancamento e DEPOIS
este script — o mesmo par de comandos para qualquer lançamento.

A base visual é o template da skill painel-dados (já com os 4 consertos de
runtime e o CSS de tabela do relatório executivo). Este script troca só o SPEC.
Célula sem medida imprime "—" (aguardando), nunca 0.
"""
import json
import re
import sys
from pathlib import Path

_V2 = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(_V2))
TEMPLATE = _V2 / ".claude" / "skills" / "painel-dados" / "template.html"
ROTULO_GOOGLE = "Google Ads"   # mesmo literal de lancamento_unidades (o contrato ja vem rotulado)
ROTULO_LEAD = "Lead Padrão (Meta)"
# baldes que NÃO são campanha de modelo (pro corte ML vs Lead da seção de público)
_NAO_MODELO = {ROTULO_LEAD, ROTULO_GOOGLE, "Orgânico/Outro", "Não está na base"}

# classificador canônico de público (fallback pra contrato antigo sem a coluna)
from src.core.ab_arm import temperatura_da_campanha  # noqa: E402


def br(v, dec=2, prefixo=""):
    """Número no formato pt-BR; None/ausente vira o traço de 'aguardando'."""
    if v is None:
        return "—"
    s = f"{v:,.{dec}f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"{prefixo}{s}"


def pct(v, dec=1):
    return "—" if v is None else f"{v:.{dec}f}%".replace(".", ",")


def _cor(v, dec=2, prefixo="R$ "):
    """Célula de dinheiro COLORIDA: verde positivo, vermelho negativo, traço ausente."""
    if v is None:
        return "—"
    return (br(v, dec, prefixo), "pos" if v > 0 else "neg")


def _tab(headers, rows, aggr=None):
    """Tabela no CSS .tb do relatório executivo (scroll horizontal próprio)."""
    th = "".join(f"<th>{h}</th>" for h in headers)
    def _tr(cells, cls=""):
        tds = "".join(
            f"<td class='{c[1]}'>{c[0]}</td>" if isinstance(c, tuple) else f"<td>{c}</td>"
            for c in cells)
        return f"<tr class='{cls}'>{tds}</tr>"
    body = "".join(_tr(r) for r in rows)
    ag = _tr(aggr, "ag") if aggr else ""
    return (f"<div class='tw'><table class='tb'><thead><tr>{th}</tr></thead>"
            f"<tbody>{body}{ag}</tbody></table></div>")


def _serie_ml_vs_lead(pasta_alvo):
    """Placar dinâmico do lucro POR CADASTRO (ML frio vs Lead frio) sobre todos
    os contratos disponíveis. Devolve (ganha, total, alvo_ganha_parcial|None)."""
    base = Path("docs/relatorios")
    fontes = {}
    for q in sorted(base.glob("_corrida/*/contrato.json")):
        fontes[q.parent.name.upper()] = q
    for q in sorted(base.glob("lf*_resultado/contrato.json")):
        fontes[q.parent.name[:-10].upper()] = q      # pasta vence a corrida
    alvo_nome = pasta_alvo.name[:-10].upper()

    def _placar_de(q):
        c = json.loads(q.read_text())
        ca = c["tabelas"]["campanhas"]
        meta_rows = [x for x in ca if x.get("plataforma") == "meta"]
        frio = [x for x in meta_rows
                if temperatura_da_campanha(x["campanha"]) != "quente"]
        ml = [x for x in frio if x["modelo"] not in _NAO_MODELO]
        ld = [x for x in frio if x["modelo"] == ROTULO_LEAD]

        def _lc(rows):
            l = sum(x.get("leads") or 0 for x in rows)
            return (sum(x.get("lucro") or 0 for x in rows) / l) if l else None
        a, b = _lc(ml), _lc(ld)
        fechado = (c["meta"]["estado"] in ("venda_fechada_imatura", "maduro")
                   and str(c["meta"].get("sales_max"))[:10] >= c["meta"]["vendas_end"][:10])
        return a, b, fechado

    ganha = total = 0
    alvo_parcial = None
    for lf, q in fontes.items():
        try:
            a, b, fechado = _placar_de(q)
        except Exception:
            continue
        if a is None or b is None:
            continue
        if lf == alvo_nome and not fechado:
            alvo_parcial = a > b
            continue
        if not fechado:
            continue
        total += 1
        ganha += 1 if a > b else 0
    return ganha, total, alvo_parcial


def main() -> int:
    pasta = Path(sys.argv[1] if len(sys.argv) > 1 else "docs/relatorios/lf64_resultado")
    c = json.loads((pasta / "contrato.json").read_text())
    m, cob, j = c["meta"], c["meta"]["cobertura"], c["julgamento"]
    ca, u, ct = (c["tabelas"]["campanhas"], c["tabelas"]["unidades"],
                 c["tabelas"]["criativos_por_tipo"])
    lf = c["lf"]
    tem_venda = bool(m["tem_venda"])
    provisorio = ("" if m["estado"] == "maduro"
                  else " · <b>PROVISÓRIO</b>" if tem_venda else "")

    tot_gasto = sum(x["gasto"] or 0 for x in ca)
    tot_cad = cob["cadastros"]
    cpl_geral = tot_gasto / tot_cad if tot_cad else None

    # ── agregado por modelo (a base da tela 1) ───────────────────────────────
    mods = {}
    for x in ca:
        d = mods.setdefault(x["modelo"], dict(g=0.0, l=0, n=0, v=None, f=None, lu=None))
        d["g"] += x["gasto"] or 0
        d["l"] += x["leads"] or 0
        d["n"] += 1
        if tem_venda:
            for k, kk in (("v", "vendas"), ("f", "faturamento"), ("lu", "lucro")):
                if x.get(kk) is not None:
                    d[k] = (d[k] or 0) + x[kk]
    ordem = sorted(mods.items(), key=lambda kv: -kv[1]["g"])

    linhas_mod = []
    for nome, d in ordem:
        cpl = d["g"] / d["l"] if d["l"] else None
        roas = (d["f"] / d["g"]) if (d["f"] is not None and d["g"]) else None
        conv = (d["v"] / d["l"]) if (d["v"] is not None and d["l"]) else None
        linhas_mod.append([f"<b>{nome}</b>", d["n"], br(d["g"], 2, "R$ "), f"{d['l']:,}".replace(",", "."),
                           br(cpl, 2, "R$ "), br(d["v"], 0),
                           pct(100 * conv, 2) if conv is not None else "—",
                           br(d["f"], 2, "R$ "),
                           br(roas, 2),
                           (br(d["lu"], 2, "R$ "), "pos" if (d["lu"] or 0) > 0 else "neg") if d["lu"] is not None else "—"])
    tot_l = sum(d["l"] for _, d in ordem)
    tot_v = sum(d["v"] or 0 for _, d in ordem) if tem_venda else None
    tot_f = sum(d["f"] or 0 for _, d in ordem) if tem_venda else None
    linhas_mod.append(["<b>TOTAL</b>", sum(d["n"] for _, d in ordem), br(tot_gasto, 2, "R$ "),
                       f"{tot_l:,}".replace(",", "."), br(cpl_geral, 2, "R$ "),
                       br(tot_v, 0) if tem_venda else "—",
                       pct(100 * tot_v / tot_l, 2) if (tem_venda and tot_l) else "—",
                       br(tot_f, 2, "R$ ") if tem_venda else "—",
                       # TOTAL soma de verdade (pedido do Ramon, 31/08): ROAS =
                       # tudo que entrou / tudo que saiu; lucro = a diferença.
                       # Inclui orgânico e 'não está na base' no faturamento.
                       br(tot_f / tot_gasto, 2) if (tem_venda and tot_gasto) else "—",
                       _cor(tot_f - tot_gasto) if tem_venda else "—"])
    tab_modelo = _tab(["Tipo de campanha", "Camp.", "Gasto", "Cadastros", "CPL",
                       "Vendas", "Conversão", "Faturamento", "ROAS", "Lucro"],
                      linhas_mod[:-1], linhas_mod[-1])

    # ── tela 1b: público quente vs frio + lucro ML vs Lead — seções FIXAS ────
    # (decisão do Ramon, 31/08: a descoberta do DEV21 — dentro do frio o modelo
    #  não separava — e a pergunta "ML lucra mais que Lead?" deixam de ser
    #  análise avulsa e entram em todo painel, todo lançamento.)
    def _temp(x):
        if x.get("plataforma") != "meta":
            return None
        t = x.get("temperatura")   # contrato novo traz a coluna; antigo não
        return t if t is not None else temperatura_da_campanha(x.get("campanha"))

    def _agg_pub(rows, rot):
        g = sum(x.get("gasto") or 0 for x in rows)
        l = sum(x.get("leads") or 0 for x in rows)
        v = sum(x.get("vendas") or 0 for x in rows) if tem_venda else None
        f = sum(x.get("faturamento") or 0 for x in rows) if tem_venda else None
        lu = sum(x.get("lucro") or 0 for x in rows) if tem_venda else None
        return [f"<b>{rot}</b>", len(rows), br(g, 2, "R$ "),
                f"{l:,}".replace(",", "."), br(g / l if l else None, 2, "R$ "),
                br(v, 0), pct(100 * v / l, 2) if (tem_venda and l) else "—",
                br(f, 2, "R$ "), br(f / g, 2) if (tem_venda and g) else "—",
                _cor(lu), _cor((lu / l) if (tem_venda and l and lu is not None) else None)]

    head_pub = ["", "Camp.", "Gasto", "Cadastros", "CPL", "Vendas", "Conversão",
                "Faturamento", "ROAS", "Lucro", "Lucro/cadastro"]
    meta_rows = [x for x in ca if x.get("plataforma") == "meta"]
    por_pub = {}
    for x in meta_rows:
        por_pub.setdefault(_temp(x), []).append(x)
    ordem_pub = sorted(por_pub, key=lambda t: -sum(y.get("gasto") or 0 for y in por_pub[t]))
    # A tabela por público só informa quando há MAIS de um público (nota do
    # Ramon, 31/08): com só frio ela repete o total e sai do painel.
    tab_pub = (_tab(head_pub, [_agg_pub(por_pub[t], t) for t in ordem_pub])
               if len(ordem_pub) >= 2 else "")

    frio_rows = [x for x in meta_rows if _temp(x) != "quente"]
    ml_frio = [x for x in frio_rows if x["modelo"] not in _NAO_MODELO]
    lead_frio = [x for x in frio_rows if x["modelo"] == ROTULO_LEAD]
    quentes = [x for x in meta_rows if _temp(x) == "quente"]
    linhas_ml = [_agg_pub(ml_frio, "Campanhas de MODELO (Champion+Challenger), frio"),
                 _agg_pub(lead_frio, "Lead Padrão (Meta), frio")]
    if quentes:
        linhas_ml.append(_agg_pub(quentes, "Público quente (economia própria, fora da disputa)"))
    tab_ml = ("<p class='h2sub' style='margin-top:14px'><b>Lucro: modelo contra Lead padrão</b> "
              "— só público frio dos dois lados, porque o quente tem economia própria "
              "(regra de 16/08) e inflaria o lado em que caísse:</p>") + _tab(head_pub, linhas_ml)
    if tem_venda and ml_frio and lead_frio:
        def _lucro_cad(rows):
            l = sum(x.get("leads") or 0 for x in rows)
            lu = sum(x.get("lucro") or 0 for x in rows)
            return (lu / l) if l else None
        a_, b_ = _lucro_cad(ml_frio), _lucro_cad(lead_frio)
        if a_ is not None and b_ is not None:
            tab_ml += (f"<p class='h2sub' style='margin-top:10px'>Neste lançamento, cada "
                       f"cadastro frio do modelo rendeu <b>{br(a_, 2, 'R$ ')}</b> de lucro; "
                       f"o do Lead padrão, <b>{br(b_, 2, 'R$ ')}</b>.")
            try:
                g_, t_, alvo_p = _serie_ml_vs_lead(pasta)
                extra = ""
                if alvo_p is not None:
                    extra = (" Incluindo este parcial, "
                             f"{g_ + (1 if alvo_p else 0)} de {t_ + 1}.")
                tab_ml += (f"<p class='h2sub' style='margin-top:6px'>Série dos lançamentos "
                           f"fechados (régua atual): o modelo ganhou em {g_} de {t_}."
                           + extra + "</p>")
            except Exception:
                pass

    sep = c.get("separacao_temperatura") or []
    tab_sep = ""
    if sep:
        rows_s = []
        for s in sep:
            lift = s.get("lift")
            rows_s.append([f"<b>{s['temperatura']}</b>",
                           f"{s['leads']:,}".replace(",", "."),
                           pct(s.get("pct_topo"), 1),
                           pct(100 * s["taxa_topo"], 2) if s.get("taxa_topo") is not None else "—",
                           pct(100 * s["taxa_base"], 2) if s.get("taxa_base") is not None else "—",
                           (br(lift, 2) + "x") if lift is not None else "—"])
        regua_sep = next((x.get("regua") for x in sep if x.get("regua")), None)
        rot_regua = (" Régua ÚNICA do <b>Champion</b>: a nota do abr_28 para todo "
                     "lead, inclusive os atendidos pelo Challenger (a nota mista "
                     "dilui a medição — decisão de 01/09)."
                     if regua_sep == "decil_champion" else "")
        tab_sep = ("<p class='h2sub' style='margin-top:14px'><b>Separação dentro de cada "
                   "público</b> (respondentes com nota; topo = notas 9-10, base = 1-8; "
                   "lift = quantas vezes o topo converte acima da base)." + rot_regua + "</p>"
                   + _tab(["Público", "Leads c/ nota", "% no topo", "Conv. topo",
                           "Conv. base", "Lift topo/base"], rows_s))

    # ── tela 2: campanha a campanha (top por gasto) ──────────────────────────
    top_camp = sorted([x for x in ca if x.get("gasto")], key=lambda x: -x["gasto"])[:15]
    linhas_c = []
    for x in top_camp:
        linhas_c.append([x["campanha"][:46], x["modelo"], br(x["gasto"], 2, "R$ "),
                         f"{x['leads']:,}".replace(",", "."), br(x["cpl"], 2, "R$ "),
                         br(x.get("vendas"), 0), br(x.get("faturamento"), 2, "R$ "),
                         br(x.get("roas"), 2), _cor(x.get("lucro"))])
    tab_camp = _tab(["Campanha", "Tipo", "Gasto", "Cadastros", "CPL", "Vendas",
                     "Faturamento", "ROAS", "Lucro"], linhas_c)
    # Destaques gerados do próprio contrato: o criativo que DOMINA a verba da
    # campanha separa positiva de negativa (pedido do Ramon, 24/08).
    if tem_venda:
        g_cria_cid = {}
        for x in u:
            if x.get("gasto"):
                d = g_cria_cid.setdefault(x["cid"], {})
                d[x["criativo"]] = d.get(x["criativo"], 0) + x["gasto"]
        pos, neg, dom_pos = [], [], {}
        for x in ca:
            if not x.get("gasto") or x["gasto"] < 1000 or x.get("lucro") is None:
                continue
            d = g_cria_cid.get(x["cid"])
            if not d:
                continue
            top = max(d, key=d.get)
            (pos if x["lucro"] > 0 else neg).append((x, top))
            if x["lucro"] > 0:
                dom_pos[top] = dom_pos.get(top, 0) + 1
        if pos and dom_pos:
            lider = max(dom_pos, key=dom_pos.get)
            # conversão do líder por tipo de público (unidades dele, por cid)
            tipo_por_cid = {x["cid"]: x["modelo"] for x in ca}
            cv = {}
            for x in u:
                if x["criativo"] != lider:
                    continue
                t = tipo_por_cid.get(x["cid"])
                if not t:
                    continue
                d = cv.setdefault(t, [0.0, 0.0])
                d[0] += x.get("cadastros") or 0
                d[1] += x.get("vendas") or 0
            # contraste = os 2 tipos onde o líder tem MAIS volume (leitura sólida)
            top2 = sorted((t for t, (l, v) in cv.items() if l >= 100 and v),
                          key=lambda t: -cv[t][0])[:2]
            cv = {t: 100 * cv[t][1] / cv[t][0] for t in top2}
            # tipo da separação = onde a fração de positivas com o líder à frente
            # mais se descola da fração de negativas (>= 3 campanhas julgadas)
            nit = {}
            for t in {x["modelo"] for x, _ in pos + neg}:
                b_ = sum(1 for x, _ in pos if x["modelo"] == t)
                d_ = sum(1 for x, _ in neg if x["modelo"] == t)
                if b_ + d_ < 3 or not b_:
                    continue
                a_ = sum(1 for x, top in pos if x["modelo"] == t and top == lider)
                c_ = sum(1 for x, top in neg if x["modelo"] == t and top == lider)
                nit[t] = a_ / b_ - (c_ / d_ if d_ else 0)
            if len(cv) >= 2 and nit:
                hi = max(cv, key=cv.get)
                lo = min(cv, key=cv.get)
                sep = max(nit, key=nit.get)
                a2 = sum(1 for x, top in pos if x["modelo"] == sep and top == lider)
                b2 = sum(1 for x, _ in pos if x["modelo"] == sep)
                c2 = sum(1 for x, top in neg if x["modelo"] == sep and top == lider)
                d2 = sum(1 for x, _ in neg if x["modelo"] == sep)
                contras = [x for x, top in neg if top == lider]
                contra_txt = ""
                if contras:
                    x0 = contras[0]
                    contra_txt = (
                        f" O porém que fecha o raciocínio: a campanha "
                        f"<b>{x0['campanha'][:34]}…</b> tinha o {lider} à frente e mesmo "
                        f"assim ficou negativa ({br(x0['lucro'], 0, 'R$ ')}), porque o "
                        f"lote de leads que ELA comprou veio fraco: criativo não salva "
                        f"público ruim.")
                tab_camp += (
                    f"<p class='h2sub' style='margin-top:10px'><b>O que tirar dessa "
                    f"tabela — por que uma campanha de ML lucra e a outra dá prejuízo?</b> "
                    f"As campanhas de ML compram gente parecida (quem escolhe o público "
                    f"é o modelo). O que separa o sinal é QUAL criativo ficou com a "
                    f"verba, e o preço pago por lead contra o teto. A prova nos "
                    f"números: pegue o <b>{lider}</b> e olhe ele em dois públicos. "
                    f"No {hi}, cada 100 cadastros dele viram ~{br(cv[hi], 2)} vendas; no "
                    f"{lo}, ~{br(cv[lo], 2)}. Mesmo anúncio, valor por lead "
                    f"{br(cv[hi] / cv[lo], 1)}x diferente: isso é o PÚBLICO definindo o "
                    f"patamar, e quem escolhe o público de cada campanha é o modelo. "
                    f"Agora trave o público e olhe só o {sep}: são {b2 + d2} campanhas "
                    f"comprando o mesmo tipo de gente; das {b2} que lucraram, {a2} têm o "
                    f"{lider} como maior verba, contra {c2} das {d2} que perderam. Mesmo "
                    f"público, criativos diferentes, sinais opostos: dentro do tipo, quem "
                    f"separa é o criativo.{contra_txt} Resposta curta: ML com a verba "
                    f"concentrada no criativo certo e CPL dentro do teto lucra; ML "
                    f"pagando lead caro em criativo fraco vira prejuízo. Como nenhuma "
                    f"das metades decide sozinha, a régua que junta as duas é o teto "
                    f"(público × criativo × preço): por isso ele julga o PAR.</p>")

    # ── tela 3: criativos por tipo (verba + qualidade + efeito na conversão) ─
    # Efeito do criativo = quanto o HISTÓRICO dele multiplica a conversão prevista
    # da unidade: peso×lift + (1−peso). Agregado por (tipo, criativo) e por tipo,
    # ponderado pelos leads — responde "quanto os criativos influenciaram a % de
    # conversão da campanha" no lado PREVISTO (o realizado é a coluna Conversão).
    rot_cid = dict(zip((x["cid"] for x in ca), (x["modelo"] for x in ca)))

    def _tipo(x):
        return ROTULO_GOOGLE if x["canal"] == "google" else (rot_cid.get(x["cid"]) or "—")

    efeito_cria, efeito_tipo, teto_cria = {}, {}, {}
    for x in u:
        mod = ROTULO_GOOGLE if x["canal"] == "google" else rot_cid.get(x["cid"])
        if mod is None:
            continue
        peso = x.get("peso_historico") or 0.0
        lift = x.get("lift_criativo")
        mult = (peso * lift + (1 - peso)) if lift is not None else 1.0
        n = int(x.get("leads_ledger") or 0)
        if x.get("teto") is not None and n:
            t = teto_cria.setdefault((mod, x["criativo"]), [0.0, 0])
            t[0] += x["teto"] * n
            t[1] += n
        for chave, d in ((("c", mod, x["criativo"]), efeito_cria),
                         (("t", mod), efeito_tipo)):
            s = d.setdefault(chave, [0.0, 0])
            s[0] += mult * n
            s[1] += n
    def _efeito(d, *k):
        s = d.get(k)
        return (s[0] / s[1] - 1.0) if s and s[1] else None

    # (tela 3 fundida na 7 em 24/08: teto e efeito viram colunas da agregada)

    # ── tela 4: teto (dentro vs acima) ───────────────────────────────────────
    modo = "MEDIDO" if tem_venda else "PREVISÃO — o julgamento fecha quando as vendas caírem"
    rows_j = []
    for rot, lado in (("Respeitaram o teto", "dentro"), ("Estouraram o teto", "acima")):
        b = j[lado]
        rows_j.append([f"<b>{rot}</b>", b["n"], br(b["gasto"], 2, "R$ "),
                       f"{b['leads']:,}".replace(",", "."), br(b.get("vendas"), 0),
                       br(b.get("roas_ponderado"), 2), _cor(b.get("lucro")),
                       br(b.get("bateu_meta"), 0), br(b.get("lucro_acima_1000"), 0),
                       br(b.get("roas_positivo"), 0)])
    tab_teto = _tab(["", "Unidades", "Gasto", "Leads", "Vendas", "ROAS",
                     "Lucro", "Bateu a meta*", "Lucro > R$ 1.000", "Deu lucro"], rows_j)
    # A régua alternativa pedida (meta 1,5): o teto publicado usa alvo 2,0, então
    # meta 1,5 equivale a teto ×(2,0/1,5). Mesmas unidades julgáveis.
    uj = [x for x in u if x.get("no_corte") and x.get("dentro_do_teto") is not None]
    d15 = [x for x in uj if x["cpl"] <= x["teto"] * (2.0 / 1.5)]
    a15 = [x for x in uj if not (x["cpl"] <= x["teto"] * (2.0 / 1.5))]

    def _lado15(xs, rot):
        gasto = sum(x.get("gasto") or 0 for x in xs)
        leads = sum(int(x.get("leads_ledger") or 0) for x in xs)
        if tem_venda and xs:
            fat = sum(x.get("faturamento") or 0 for x in xs)
            vendas = sum(x.get("vendas") or 0 for x in xs)
            lucro = sum(x.get("lucro") or 0 for x in xs)
            roas = fat / gasto if gasto else None
            bateu = sum(1 for x in xs if (x.get("roas") or 0) >= 1.5 * (1 - j["tolerancia_meta"]))
            l1000 = sum(1 for x in xs if (x.get("lucro") or 0) > 1000)
            pos = sum(1 for x in xs if (x.get("lucro") or 0) > 0)
        else:
            vendas = roas = lucro = bateu = l1000 = pos = None
        return [f"<b>{rot}</b>", len(xs), br(gasto, 2, "R$ "),
                f"{leads:,}".replace(",", "."), br(vendas, 0), br(roas, 2),
                _cor(lucro), br(bateu, 0), br(l1000, 0), br(pos, 0)]

    tab_teto15 = _tab(["", "Unidades", "Gasto", "Leads", "Vendas", "ROAS",
                       "Lucro", "Bateu a meta†", "Lucro > R$ 1.000", "Deu lucro"],
                      [_lado15(d15, "Respeitaram o teto"), _lado15(a15, "Estouraram o teto")])
    tab_teto += ("<p class='h2sub' style='margin-top:14px'>A mesma leitura com meta de ROAS "
                 "<b>1,5</b> (o teto cresce ×1,33; †bateu a meta = ROAS ≥ 1,5 com a mesma "
                 "tolerância de 2%):</p>" + tab_teto15
                 + f"<p class='h2sub' style='margin-top:10px'>Com a régua mais folgada, "
                 f"<b>{len(d15)}</b> de {len(uj)} unidades couberam no teto.</p>")

    # ── o dinheiro que o corte NÃO julga (decisão do Ramon, 01/09) ───────────
    # O corte exige 100 leads e R$ 300 na dupla; CPL alto atrasa os 100 leads,
    # então o pior dinheiro fica invisível. Backtest de 01/09 (10 LFs fechados):
    # liberar o teto do PAR já com R$ 300 separa lucro de prejuízo.
    um_meta = [x for x in u if x.get("canal") == "meta" and x.get("gasto")]
    g_meta_u = sum(x["gasto"] for x in um_meta)
    sem_ver = [x for x in um_meta if not x.get("no_corte")]
    g_sem = sum(x["gasto"] for x in sem_ver)
    delta_sv = sorted([x for x in sem_ver if x["gasto"] >= 300
                       and int(x.get("leads_ledger") or 0) < 100
                       and x.get("cpl") is not None and x.get("teto") is not None],
                      key=lambda x: -x["gasto"])
    if tem_venda and delta_sv and g_meta_u:
        rows_sv = []
        for x in delta_sv[:10]:
            t15 = x["teto"] * (2.0 / 1.5)
            raz = x["cpl"] / t15
            rows_sv.append([x["criativo"][:24], str(x.get("campanha"))[:30],
                            br(x["gasto"], 0, "R$ "), int(x.get("leads_ledger") or 0),
                            br(x["cpl"]), br(t15),
                            ((br(raz, 1) + "x", "neg") if raz > 1 else ("dentro", "pos"))])
        g_delta_sv = sum(x["gasto"] for x in delta_sv)
        n_est = sum(1 for x in delta_sv if x["cpl"] > x["teto"] * (2.0 / 1.5))
        tab_teto += (
            f"<p class='h2sub' style='margin-top:14px'><b>O dinheiro que o corte não julga.</b> "
            f"O julgamento acima exige 100 leads e R$ 300 de gasto na dupla criativo×campanha. "
            f"Ficou sem veredito neste lançamento: {br(g_sem, 0, 'R$ ')} "
            f"({br(100 * g_sem / g_meta_u, 1)}% do gasto Meta). A fatia julgável ANTECIPADO "
            f"(R$ 300+ de gasto, ainda sem 100 leads) tem {len(delta_sv)} duplas e "
            f"{br(g_delta_sv, 0, 'R$ ')}, {n_est} delas estourando o teto 1,5. É o pior "
            f"dinheiro do lançamento: CPL alto é justamente o que impede de juntar 100 leads.</p>"
            + _tab(["Criativo", "Campanha", "Gasto", "Leads", "CPL", "Teto 1,5", "Estouro"],
                   rows_sv)
            + "<p class='h2sub' style='margin-top:10px'><b>Orientação nova (backtest de 01/09, "
              "10 lançamentos fechados):</b> julgar a dupla pelo teto do PRÓPRIO PAR assim que "
              "ela passa de R$ 300 de gasto separa: o dinheiro que essa regra marcou FORA rendeu "
              "0,75 por real investido; o que ela marcou DENTRO rendeu 2,66. O papel da regra é "
              "dar visibilidade a esse dinheiro no relatório, não cortar sozinha.</p>")
    # (análise autoral do porquê — ex.: o CPL 39% mais caro do LF64 — vive em
    #  conclusao.html/notas.html na pasta do LF, nunca aqui: o script é genérico
    #  e o texto do LF64 vazou pro painel do LF65 na primeira prova de reuso.)

    # ── telas 5 e 6: menor CPL e ranking ─────────────────────────────────────
    # (telas 'menor CPL' e 'ranking por lucro' removidas a pedido do Ramon, 24-25/08)


    # ── tela 7: criativo AGREGADO por tipo (todas as campanhas do tipo somadas) ──
    # No HTML só entra gasto > R$ 300 (mesmo piso do corte do teto); a lista
    # COMPLETA sai em XLSX ao lado do painel, na mesma rodada.
    chave7 = "lucro" if tem_venda else "gasto"
    ct_ord = sorted([x for x in ct if x.get("gasto")],
                    key=lambda x: -(x.get(chave7) or 0))
    def _roas7(x):
        return (x["faturamento"] / x["gasto"]) if (x.get("faturamento") is not None
                                                   and x.get("gasto")) else None
    rows7 = []
    for x in ct_ord:
        if (x.get("gasto") or 0) <= 300:
            continue
        tc = teto_cria.get((x["modelo"], x["criativo"]))
        t2 = tc[0] / tc[1] if tc and tc[1] else None
        rows7.append([x["criativo"][:38], x["modelo"], br(x.get("gasto"), 2, "R$ "),
                      br(x.get("cpl"), 2, "R$ "), br(t2, 2, "R$ "),
                      br(t2 * (2.0 / 1.5), 2, "R$ ") if t2 is not None else "—",
                      br(x.get("faturamento"), 2, "R$ "), br(_roas7(x), 2), _cor(x.get("lucro"))])
    tab_cria_tipo = _tab(["Criativo", "Tipo", "Gasto", "CPL", "Teto 2x", "Teto 1,5x",
                          "Faturamento", "ROAS", "Lucro"], rows7)
    tab_cria_tipo += (
        "<p class='h2sub' style='margin-top:10px'>Os tetos acima são o agregado do criativo "
        "no tipo (média ponderada pelos leads das campanhas em que ele rodou). Se precisarem "
        "da quebra por campanha — quanto cada campanha contribuiu para esse número, útil "
        "quando há configurações de campanha diferentes — é só pedir: a base já existe.</p>")

    # Lucrou ACIMA do teto (nota do Ramon, 31/08): ou o teto está desinformado
    # (criativo sem histórico próprio), ou a amostra é pequena, ou o estouro é
    # real e o teto merece releitura — o painel diz qual é o caso.
    if tem_venda:
        exp = []
        for x in ct_ord:
            if (x.get("gasto") or 0) <= 300 or (x.get("lucro") or 0) <= 0:
                continue
            tc = teto_cria.get((x["modelo"], x["criativo"]))
            t15 = tc[0] / tc[1] * (2.0 / 1.5) if tc and tc[1] else None
            cpl = x.get("cpl")
            if t15 is None or cpl is None or cpl <= t15:
                continue
            v = int(x.get("vendas") or 0)
            n_h = int(x.get("n_hist") or 0)
            if n_h < 1000:
                motivo = (f"histórico de só {n_h} leads — o teto quase não conhece "
                          "o criativo e sai apertado demais; estourar e lucrar é "
                          "esperado até o histórico formar")
            elif v < 3:
                motivo = (f"{v} venda(s) em {int(x.get('leads_ledger') or 0)} leads — "
                          "amostra pequena, pode ser sorte")
            else:
                motivo = (f"histórico de {n_h:,} leads e {v} vendas — estouro "
                          "lucrativo REAL; o teto dele merece releitura"
                          ).replace(",", ".")
            exp.append(f"<b>{x['criativo'][:32]}</b> ({x['modelo']}): CPL {br(cpl)} "
                       f"contra teto 1,5 de {br(t15)} — {motivo}")
        if exp:
            tab_cria_tipo += ("<p class='h2sub' style='margin-top:10px'><b>Estourou o "
                              "teto e mesmo assim lucrou — por quê?</b> "
                              + "; ".join(exp) + ".</p>")

    # ── lucro por decil (decisão do Ramon, 01/09): o custo de um lead é o CPL
    # da dupla criativo×campanha que o trouxe (a verba sai antes da nota), então
    # dá pra somar custo por decil e fechar lucro. Contrato antigo não tem o
    # bloco e a seção simplesmente não sai.
    tab_ld = ""
    ld = c["tabelas"].get("lucro_decil")
    if ld and ld.get("champion"):
        rows_ld = []
        for r in ld["champion"]["decis"]:
            rows_ld.append([f"<b>D{r['decil']}</b>", f"{r['leads']:,}".replace(",", "."),
                            br(r["custo"], 0, "R$ "), br(r["vendas"], 0),
                            br(r["faturamento"], 0, "R$ "), _cor(r["lucro"], 0),
                            _cor((r["lucro"] / r["leads"]) if r["leads"] else None)])
        resumo = []
        for nome, chave in (("Champion (abr_28)", "champion"),
                            ("Challenger (jul_24)", "challenger")):
            bloco = ld.get(chave)
            if not bloco:
                continue
            for rot, k in (("top 30 (D8-D10)", "top30"), ("resto (D1-D7)", "resto"),
                           ("fundo (D1-D5)", "fundo")):
                sm = bloco.get(k)
                if sm:
                    resumo.append([f"<b>{nome} · {rot}</b>",
                                   f"{sm['leads']:,}".replace(",", "."),
                                   br(sm["custo"], 0, "R$ "), br(sm["vendas"], 0),
                                   br(sm["faturamento"], 0, "R$ "), _cor(sm["lucro"], 0),
                                   br(sm["roas"], 2)])
        tab_ld = (
            "<p class='h2sub' style='margin-top:14px'><b>Lucro por decil</b> (só Meta frio; "
            "o custo de cada lead é o CPL da dupla criativo×campanha que o trouxe, porque a "
            "verba sai antes de o modelo dar a nota; réguas separadas por modelo, nunca "
            "misturadas). Primeiro o Champion, decil a decil:</p>"
            + _tab(["Decil", "Leads", "Custo", "Vendas", "Faturamento", "Lucro",
                    "Lucro/lead"], rows_ld)
            + "<p class='h2sub' style='margin-top:10px'>O resumo que decide, nos dois "
              "modelos (ROAS = faturamento / custo alocado):</p>"
            + _tab(["", "Leads", "Custo", "Vendas", "Faturamento", "Lucro", "ROAS"],
                   resumo))

    import pandas as pd
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    xlsx = pasta / f"criativos_por_tipo_{lf.lower()}.xlsx"
    df7 = pd.DataFrame([{
        "Criativo": x["criativo"], "Tipo": x["modelo"], "Unidades": x.get("unidades"),
        "Leads": x.get("leads_ledger"), "Cadastros": x.get("cadastros"),
        "Gasto": x.get("gasto"), "CPL": x.get("cpl"), "Vendas": x.get("vendas"),
        "Faturamento": x.get("faturamento"), "ROAS": _roas7(x), "Lucro": x.get("lucro"),
        "Conversão %": (100 * x["conversao"]) if x.get("conversao") is not None else None,
        "% notas 9-10": x.get("pct_d9_d10"),
    } for x in ct_ord])
    # Formatação no padrão do painel: dinheiro em R$ com 2 casas, ROAS 2 casas,
    # percentuais com o símbolo, inteiros sem casa; cabeçalho em negrito e congelado.
    _FMT = {"Gasto": 'R$ #,##0.00', "CPL": 'R$ #,##0.00', "Faturamento": 'R$ #,##0.00',
            "Lucro": 'R$ #,##0.00', "ROAS": '0.00', "Conversão %": '0.00"%"',
            "% notas 9-10": '0.0"%"', "Leads": '#,##0', "Cadastros": '#,##0',
            "Unidades": '0', "Vendas": '0'}
    with pd.ExcelWriter(xlsx, engine="openpyxl") as xw:
        df7.to_excel(xw, index=False, sheet_name="criativos_por_tipo")
        ws = xw.sheets["criativos_por_tipo"]
        for i, col in enumerate(df7.columns, 1):
            letra = get_column_letter(i)
            ws.column_dimensions[letra].width = (44 if col == "Criativo" else
                                                 24 if col == "Tipo" else
                                                 max(len(col) + 3, 12))
            f = _FMT.get(col)
            if f:
                for cell in ws[letra][1:]:
                    cell.number_format = f
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
    print(f"xlsx:   {xlsx}  ({len(ct_ord)} linhas, completo, formatado)")

    # ── SPEC ─────────────────────────────────────────────────────────────────
    aguardando = ("" if tem_venda else
                  "<p><i>Carrinho aberto em " + m["vendas_start"][8:10] + "/" + m["vendas_start"][5:7]
                  + ". As colunas de vendas, faturamento, ROAS e lucro preenchem sozinhas "
                  "na próxima rodada deste mesmo relatório, quando as vendas caírem — traço (—) "
                  "significa \"ainda não medido\", nunca zero.</i></p>")
    spec = {
        "eyebrow": f"DevClub · Lançamento {lf} · captação {m['cap_start'][8:10]}/{m['cap_start'][5:7]}"
                   f" a {m['cap_end'][8:10]}/{m['cap_end'][5:7]} · carrinho {m['vendas_start'][8:10]}/"
                   f"{m['vendas_start'][5:7]} a {m['vendas_end'][8:10]}/{m['vendas_end'][5:7]}",
        "title": f"Painel do {lf}",
        "subtitle": f"O mesmo relatório do DEV21, agora por máquina: um comando gera, o mesmo comando "
                    f"atualiza. Estado: <b>{m['estado'].replace('_', ' ')}</b>{provisorio}.",
        "source": f"Contrato congelado em {m['gerado_em'][:16]} · referência {m['referencia_id']} · "
                  f"régua {m['ruler_run_id'][:8]} (abr_28) · {br(100 * (cob['pct_gasto_casado'] or 0), 1)}% "
                  f"do gasto Meta casado por anúncio",
        "tiles": [
            {"color": "c1", "label": "Investimento total", "value": br(tot_gasto, 0, "R$ "),
             "note": f"Meta com imposto (×{br(m['meta_gross_up'], 2)}) + Google."},
            {"color": "accent", "label": "Cadastros", "value": f"{tot_cad:,}".replace(",", "."),
             "note": f"CPL geral {br(cpl_geral, 2, 'R$ ')} por cadastro."},
            {"color": "c2", "label": "Vendas",
             "value": (br(sum(x.get('vendas') or 0 for x in ca), 0) if tem_venda else "aguardando"),
             "note": ("Somatório das campanhas casadas." if tem_venda
                      else "Carrinho abriu hoje; preenche na próxima rodada.")},
        ],
        "sections": [
            {"title": "1 · Onde o dinheiro foi — por tipo de campanha",
             "sub": "Tipo = o MODELO que escolhe o público da campanha (identidade estável, "
                    "mesmo com os nomes novos da equipe de tráfego).",
             "html": tab_modelo + aguardando},
            {"title": "2 · Lucro do modelo contra o Lead — e a separação por público",
             "sub": "",
             "html": tab_pub + tab_ml + tab_sep + tab_ld},
            {"title": "2 · Campanha a campanha (top 15 por gasto)",
             "sub": "As mesmas colunas do debriefing: gasto, cadastros, CPL, vendas, ROAS e lucro.",
             "html": tab_camp},
            {"title": f"4 · Teto de CPL — {modo}",
             "sub": "Unidade = criativo×campanha com ≥R$ 300 de gasto. "
                    "*Meta = ROAS 2,0 com tolerância de 2% (um 1,97 conta).",
             "html": tab_teto},
            {"title": "7 · Criativo agregado por tipo de campanha",
             "html": tab_cria_tipo},
        ],
        "footer": f"Gerado do contrato.json por render_painel_lancamento.py · {m['gerado_em'][:16]} · "
                  "traço (—) = ainda não medido",
    }

    # Conclusão do lançamento: análise autoral deste LF vive em conclusao.html na
    # pasta do relatório (não no script, que é genérico); entra ANTES do carimbo.
    # Nota por-LF sob a tabela do criativo agregado: nota_criativos.html na pasta.
    nota_cria = pasta / "nota_criativos.html"
    if nota_cria.exists():
        sec = next(x for x in spec["sections"] if "Criativo agregado" in x["title"])
        sec["html"] += nota_cria.read_text()

    # Comparação com os lançamentos anteriores: gerada por
    # scripts/comparativo_lancamentos.py na pasta do LF; entra antes do carimbo.
    comparativo = pasta / "comparativo.html"
    if comparativo.exists():
        spec["sections"].append({"title": "Comparação com os lançamentos anteriores",
                                 "html": comparativo.read_text()})

    conclusao = pasta / "conclusao.html"
    if conclusao.exists():
        spec["sections"].append({"title": "Ações e recomendações",
                                 "html": conclusao.read_text()})
    for i, x in enumerate(spec["sections"], 1):   # renumera 1..N (sempre)
        x["title"] = re.sub(r"^\d+ · ", "", x["title"])
        x["title"] = f"{i} · {x['title']}"

    # Notas do lançamento: fatos pontuais deste LF vivem em notas.html na pasta
    # do relatório (não no script, que é genérico); se existir, vira a seção final.
    notas = pasta / "notas.html"
    if notas.exists():
        spec["sections"].append({"title": f"{len(spec['sections']) + 1} · Notas do lançamento",
                                 "html": notas.read_text()})

    tpl = TEMPLATE.read_text()
    ini = tpl.index("const SPEC = {")
    fim = tpl.index("/* ==== runtime")
    html = tpl[:ini] + "const SPEC = " + json.dumps(spec, ensure_ascii=False, indent=1) + ";\n\n" + tpl[fim:]
    html = re.sub(r"<title>.*?</title>", f"<title>Painel do {lf}</title>", html)
    dst = pasta / f"painel_{lf.lower()}.html"
    dst.write_text(html)
    print(f"painel: {dst}  ({dst.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
