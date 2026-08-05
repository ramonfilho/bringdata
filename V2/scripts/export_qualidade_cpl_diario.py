#!/usr/bin/env python3
"""Relatorio diario: qualidade media (decil Challenger) + CPL de alta qualidade
(D9-D10) por plataforma, nos ultimos N dias. Uma linha por dia BRT.

Read-only. Reusa as funcoes de dados que ja existem no projeto:
  - registros_ml (ledger Cloud SQL) via open_ledger_read_connection
        -> nota media (AVG decil_challenger) + contagem D9-D10 por fonte
  - MetaAPIClient.get_daily_campaign_metrics
        -> spend Meta de captacao (campanhas 'CAP') por dia, ja em BRL
  - GoogleAdsReportingClient.get_campaign_metrics
        -> spend Google por dia (cost_micros -> BRL), por dia
  - core.launches.resolve_active_launch_brt
        -> LF (lancamento) de cada dia

CPL por plataforma = spend da plataforma no dia / leads D9-D10 daquela fonte no
dia. Denominador = so alta qualidade (decil 9 ou 10), decisao do usuario.

A nota media NAO filtra por challenger_run_id (igual ao launch_score_geral atual):
a coluna decil_challenger foi backfillada por copia no ledger e nem toda linha
antiga carrega o run_id, entao filtrar subcontaria os dias antes de 08/07.

Saida: V2/outputs/qualidade_cpl_diario_<timestamp>.xlsx

Uso:
    python -m scripts.export_qualidade_cpl_diario                 # 45 dias
    python -m scripts.export_qualidade_cpl_diario --days 3        # smoke
    python -m scripts.export_qualidade_cpl_diario --end 2026-07-15 --days 45
"""
from __future__ import annotations

import argparse
import os
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]  # .../V2
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

try:
    from dotenv import load_dotenv
    load_dotenv(ROOT / ".env")
except Exception:
    pass

# Fontes que classificam a origem do lead no painel de decis (mesma regra do app.py).
SRC_META = frozenset({"facebook-ads", "fb", "ig"})
SRC_GGL = frozenset({"google-ads"})
BRT_OFFSET_H = 3  # BRT = UTC-3; created_at no ledger e UTC (naive)


def _brt_day_start_utc(d: date) -> datetime:
    """00:00 BRT de `d` -> UTC naive (para comparar com created_at do ledger)."""
    return datetime(d.year, d.month, d.day, BRT_OFFSET_H, 0, 0)


# ─────────────────────────── 1. ledger: nota media + D9-D10 por fonte ──────────
# Dedup por (email, dia BRT): cada lead conta UMA vez em cada dia em que aparece.
# NAO deduplicar global por email sobre a janela toda -> isso jogaria o lead pro
# dia da ultima linha dele e roubaria leads dos dias antigos (contagem instavel,
# depende do tamanho da janela). A chave (email, dia) da contagem diaria estavel.
_LEDGER_SQL = """
WITH per_lead AS (
  SELECT DISTINCT ON (lower(email), (created_at - INTERVAL '3 hours')::date)
         (created_at - INTERVAL '3 hours')::date AS dia_brt,
         lower(utm_source)                       AS src,
         decil_challenger                        AS decil
  FROM registros_ml
  WHERE created_at >= :ws AND created_at < :we
    AND decil_challenger IS NOT NULL
  ORDER BY lower(email), (created_at - INTERVAL '3 hours')::date,
           (utm_source IS NOT NULL AND utm_source <> '') DESC, created_at DESC
)
SELECT dia_brt, src, decil FROM per_lead
"""


def fetch_ledger_daily(start_d: date, end_d: date) -> dict:
    """{iso_date: {'decis':[...], 'meta_d9d10':n, 'ggl_d9d10':n,
                   'meta_leads':n, 'ggl_leads':n}} por dia BRT."""
    from src.data.ledger_connection import open_ledger_read_connection

    ws = _brt_day_start_utc(start_d)
    we = _brt_day_start_utc(end_d + timedelta(days=1))  # fim exclusivo
    conn = open_ledger_read_connection()
    try:
        rows = conn.run(_LEDGER_SQL, ws=ws, we=we)
    finally:
        try:
            conn.close()
        except Exception:
            pass

    out: dict = defaultdict(
        lambda: {"decis": [], "meta_d9d10": 0, "ggl_d9d10": 0,
                 "meta_leads": 0, "ggl_leads": 0}
    )
    for dia, src, decil in rows:
        iso = dia.isoformat() if hasattr(dia, "isoformat") else str(dia)
        rec = out[iso]
        d = int(decil)
        rec["decis"].append(d)
        if src in SRC_META:
            rec["meta_leads"] += 1
            if d >= 9:
                rec["meta_d9d10"] += 1
        elif src in SRC_GGL:
            rec["ggl_leads"] += 1
            if d >= 9:
                rec["ggl_d9d10"] += 1
    return out


# ─────────────────────────── 2. Meta: spend de captacao por dia ────────────────
def fetch_meta_spend_daily(start_d: date, end_d: date) -> dict:
    """{iso_date: spend_brl} das campanhas de captacao (nome contem 'CAP')."""
    from src.validation.meta_api_client import MetaAPIClient

    mc = MetaAPIClient()
    df = mc.get_daily_campaign_metrics(
        start_d.isoformat(), end_d.isoformat(), apply_filters=True
    )
    spend: dict = defaultdict(float)
    if df is None or len(df) == 0:
        return spend
    for _, r in df.iterrows():
        dstr = str(r.get("date") or "").strip()
        if not dstr:
            continue
        spend[dstr] += float(r.get("spend_dia") or 0)
    return spend


# ─────────────────────────── 3. Google: spend por dia ─────────────────────────
def fetch_google_spend_daily(days: list[date]) -> dict:
    """{iso_date: spend_brl | None}. None = falha na API pra aquele dia.

    statuses=None inclui campanhas hoje pausadas (o gasto historico continua
    valendo; filtrar por status ENABLED atual perderia dias antigos)."""
    from src.validation.google_ads_api_client import GoogleAdsReportingClient

    cust = os.environ["GOOGLE_ADS_CUSTOMER_ID"].replace("-", "")
    gc = GoogleAdsReportingClient(cust)
    spend: dict = {}
    for d in days:
        ds = d.isoformat()
        try:
            rows = gc.get_campaign_metrics(ds, ds, statuses=None)
            spend[ds] = round(sum(float(x.get("spend") or 0) for x in rows), 2)
        except Exception as e:
            print(f"  ! Google {ds}: {e}", file=sys.stderr)
            spend[ds] = None
    return spend


# ─────────────────────────── 4. dia -> LF ─────────────────────────────────────
def lf_for_day(d: date) -> str:
    from src.core.launches import resolve_active_launch_brt

    active = resolve_active_launch_brt(today=d)
    return active.name if active else ""


# ─────────────────────────── montagem + export ────────────────────────────────
def build_rows(days, ledger, meta_spend, ggl_spend):
    rows = []
    for d in days:
        iso = d.isoformat()
        L = ledger.get(iso)
        decis = L["decis"] if L else []
        n = len(decis)
        nota = round(sum(decis) / n, 2) if n else None
        n_d9d10 = sum(1 for x in decis if x >= 9)
        pct_d9d10 = round(n_d9d10 / n * 100, 1) if n else None

        m_leads = L["meta_d9d10"] if L else 0
        g_leads = L["ggl_d9d10"] if L else 0
        m_sp = meta_spend.get(iso)
        g_sp = ggl_spend.get(iso)
        m_cpl = round(m_sp / m_leads, 2) if (m_sp is not None and m_leads > 0) else None
        g_cpl = round(g_sp / g_leads, 2) if (g_sp is not None and g_leads > 0) else None

        rows.append({
            "Data": d.strftime("%d/%m/%Y"),
            "LF": lf_for_day(d),
            "Nota media (Challenger)": nota,
            "% D9-D10 (geral)": pct_d9d10,
            "Leads (geral)": n,
            "Leads D9-D10 Meta": m_leads,
            "Spend Meta (R$)": round(m_sp, 2) if m_sp is not None else None,
            "CPL Meta D9-D10 (R$)": m_cpl,
            "Leads D9-D10 Google": g_leads,
            "Spend Google (R$)": g_sp,
            "CPL Google D9-D10 (R$)": g_cpl,
        })
    return rows


def export_xlsx(rows, output_path: Path):
    import pandas as pd
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    df = pd.DataFrame(rows)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    money_cols = {"Spend Meta (R$)", "Spend Google (R$)",
                  "CPL Meta D9-D10 (R$)", "CPL Google D9-D10 (R$)"}
    num2_cols = {"Nota media (Challenger)"}
    pct_cols = {"% D9-D10 (geral)"}

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Diario")
        ws = writer.sheets["Diario"]

        widths = {
            "Data": 12, "LF": 8, "Nota media (Challenger)": 22,
            "% D9-D10 (geral)": 16, "Leads (geral)": 12,
            "Leads D9-D10 Meta": 18, "Spend Meta (R$)": 16,
            "CPL Meta D9-D10 (R$)": 20, "Leads D9-D10 Google": 20,
            "Spend Google (R$)": 18, "CPL Google D9-D10 (R$)": 22,
        }
        for i, col in enumerate(df.columns, start=1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(col, 14)

        header_fill = PatternFill("solid", fgColor="1E3A5F")
        white = Font(bold=True, color="FFFFFF")
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = white
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32

        col_idx = {col: i + 1 for i, col in enumerate(df.columns)}
        for col in money_cols:
            letter = get_column_letter(col_idx[col])
            for row in range(2, len(df) + 2):
                ws[f"{letter}{row}"].number_format = 'R$ #,##0.00'
        for col in num2_cols:
            letter = get_column_letter(col_idx[col])
            for row in range(2, len(df) + 2):
                ws[f"{letter}{row}"].number_format = '0.00'
        for col in pct_cols:
            letter = get_column_letter(col_idx[col])
            for row in range(2, len(df) + 2):
                ws[f"{letter}{row}"].number_format = '0.0"%"'

        ws.freeze_panes = "A2"

    return df


def main():
    ap = argparse.ArgumentParser(
        description="Relatorio diario de qualidade (decil Challenger) + CPL D9-D10 por plataforma."
    )
    ap.add_argument("--days", type=int, default=45, help="janela em dias (default 45)")
    ap.add_argument("--end", default=None, metavar="YYYY-MM-DD",
                    help="ultimo dia da janela (default: hoje BRT)")
    ap.add_argument("--output", default=None, metavar="CAMINHO")
    ap.add_argument("--no-google", action="store_true",
                    help="pula o pull do Google (mais rapido no smoke)")
    args = ap.parse_args()

    if args.end:
        end_d = date.fromisoformat(args.end)
    else:
        end_d = (datetime.utcnow() - timedelta(hours=BRT_OFFSET_H)).date()
    start_d = end_d - timedelta(days=args.days - 1)
    days = [start_d + timedelta(days=i) for i in range((end_d - start_d).days + 1)]

    print(f"Janela BRT: {start_d.isoformat()} .. {end_d.isoformat()} ({len(days)} dias)")

    print("1/4 ledger (nota media + D9-D10 por fonte)...")
    ledger = fetch_ledger_daily(start_d, end_d)
    print(f"    {len(ledger)} dias com leads scoreados")

    print("2/4 Meta spend diario (campanhas CAP)...")
    meta_spend = fetch_meta_spend_daily(start_d, end_d)
    print(f"    {sum(1 for v in meta_spend.values() if v)} dias com spend Meta")

    if args.no_google:
        print("3/4 Google spend: PULADO (--no-google)")
        ggl_spend = {}
    else:
        print(f"3/4 Google spend diario ({len(days)} chamadas)...")
        ggl_spend = fetch_google_spend_daily(days)
        print(f"    {sum(1 for v in ggl_spend.values() if v)} dias com spend Google")

    print("4/4 montando + exportando...")
    rows = build_rows(days, ledger, meta_spend, ggl_spend)

    if args.output:
        output_path = Path(args.output)
    else:
        ts = (datetime.utcnow() - timedelta(hours=BRT_OFFSET_H)).strftime("%Y%m%d_%H%M")
        output_path = ROOT / "outputs" / f"qualidade_cpl_diario_{ts}.xlsx"

    df = export_xlsx(rows, output_path)

    print("\n" + "=" * 60)
    with __import__("pandas").option_context("display.max_columns", None, "display.width", 200):
        print(df.to_string(index=False))
    print("=" * 60)
    print(f"Salvo em: {output_path}")


if __name__ == "__main__":
    main()
